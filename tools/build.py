#!/usr/bin/env python3
"""
Loom2hand site builder.

Reads products.xlsx + images/<folder>/*.jpg and writes:
  site/assets/products.js   (public catalogue – no buying rates)
  site/img/*.webp           (optimised photos, 480px + 960px)

Usage:  python tools/build.py
Needs:  pip install openpyxl pillow
"""
import json, math, os, re, sys, time
from openpyxl import load_workbook
from PIL import Image, ImageOps

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, 'products.xlsx')
IMG_SRC = os.path.join(ROOT, 'images')
SITE = os.path.join(ROOT, 'site')
IMG_OUT = os.path.join(SITE, 'img')
SIZES = (480, 960)
EXTS = ('.jpg', '.jpeg', '.png', '.webp')


def num(v, default=0.0):
    try:
        if v is None or v == '':
            return default
        return float(str(v).replace('%', '').replace(',', '').replace('₹', '').strip())
    except ValueError:
        return default


def pct(v, default):
    x = num(v, default)
    return x / 100 if x > 1 else x


def read_settings(wb):
    s = {}
    for row in wb['Settings'].iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            s[str(row[0]).strip()] = row[1]
    return s


def prices(buy, markup, disc):
    """Same rule as the Excel formulas."""
    sell = math.ceil((buy * (1 + markup) + 1) / 10) * 10 - 1
    mrp = int(math.floor((sell / (1 - disc) + 1) / 10 + 0.5)) * 10 - 1 if disc > 0 else sell
    mrp = max(mrp, sell)
    off = int((mrp - sell) * 100 // mrp) if mrp else 0
    return int(sell), int(mrp), off


def label_from(fn, i):
    stem = os.path.splitext(fn)[0]
    m = re.match(r'^\d+[_\- ]?(.*)$', stem)
    lab = (m.group(1) if m else stem).replace('-', ' ').replace('_', ' ').strip()
    return lab.title() if lab else f'Design {i}'


def make_webp(src, key):
    outs = []
    st = os.path.getmtime(src)
    im = None
    for w in SIZES:
        out = os.path.join(IMG_OUT, f'{key}-{w}.webp')
        outs.append(os.path.basename(out))
        if os.path.exists(out) and os.path.getmtime(out) >= st:
            continue
        if im is None:
            im = ImageOps.exif_transpose(Image.open(src)).convert('RGB')
            # square-pad anything that isn't square, on the site's paper colour
            if im.width != im.height:
                side = max(im.size)
                bg = Image.new('RGB', (side, side), (247, 243, 236))
                bg.paste(im, ((side - im.width) // 2, (side - im.height) // 2))
                im = bg
        im.resize((w, w), Image.LANCZOS).save(out, 'WEBP', quality=74 if w == 480 else 70, method=6)
    return outs


def inch_to_cm(s):
    try:
        a, b = [float(x) for x in re.split(r'\s*[xX×]\s*', str(s))[:2]]
        return f'{a:g}" × {b:g}"', f'{round(a*2.54)} × {round(b*2.54)} cm'
    except Exception:
        return str(s or ''), ''


def main():
    if not os.path.exists(XLSX):
        sys.exit('products.xlsx not found')
    wb = load_workbook(XLSX, data_only=False)
    S = read_settings(wb)
    markup = pct(S.get('Markup %'), 0.5)
    disc = pct(S.get('Display discount %'), 0.25)
    low = int(num(S.get('Low stock alert below'), 10))
    os.makedirs(IMG_OUT, exist_ok=True)

    ws = wb['Products']
    head = [str(c.value).strip() if c.value else '' for c in ws[1]]
    col = {h: i for i, h in enumerate(head)}

    def g(row, name):
        i = col.get(name)
        return row[i] if i is not None and i < len(row) else None

    products, used, warnings = [], set(), []
    for row in ws.iter_rows(min_row=2, values_only=True):
        sku = g(row, 'SKU')
        if not sku:
            continue
        sku = str(sku).strip()
        cat = str(g(row, 'Category') or 'Towels').strip()
        active = str(g(row, 'Active (Y/N)') or '').strip().upper() == 'Y'
        buy = num(g(row, 'Buy Rate (₹)'), 0)
        if not active:
            continue
        if buy <= 0:
            warnings.append(f'{sku}: active but no Buy Rate – skipped')
            continue
        folder = str(g(row, 'Image Folder') or sku).strip()
        d = os.path.join(IMG_SRC, folder)
        files = sorted(f for f in os.listdir(d) if f.lower().endswith(EXTS)) if os.path.isdir(d) else []
        if not files:
            warnings.append(f'{sku}: no photos yet in images/{folder} – showing a drawing instead')
        imgs = []
        for i, f in enumerate(files, 1):
            key = re.sub(r'[^a-z0-9]+', '-', f'{folder}-{os.path.splitext(f)[0]}'.lower()).strip('-')
            a, b = make_webp(os.path.join(d, f), key)
            used.update((a, b))
            imgs.append({'s': a, 'l': b, 'label': label_from(f, i)})
        sell, mrp, off = prices(buy, markup, disc)
        stock = g(row, 'Stock Qty')
        stock = None if stock in (None, '') else int(num(stock))
        size_in, size_cm = inch_to_cm(g(row, 'Size (inch)'))
        products.append({
            'sku': sku, 'cat': cat, 'name': str(g(row, 'Website Name') or sku).strip(),
            'tag': str(g(row, 'Filter Tag') or '').strip(), 'weave': str(g(row, 'Weave') or '').strip(),
            'yarn': str(g(row, 'Yarn Count') or '').strip(), 'size': size_in, 'sizeCm': size_cm,
            'gsm': int(num(g(row, 'Weight (g)'))) or None, 'price': sell, 'mrp': mrp, 'off': off,
            'stock': stock, 'featured': str(g(row, 'Featured (Y/N)') or '').strip().upper() == 'Y',
            'desc': str(g(row, 'Description (optional)') or '').strip(), 'img': imgs,
            'art': ('a-kerchief' if cat == 'Napkins' else 'a-jute' if cat == 'Jute Bags' else 'a-white' if str(g(row, 'Filter Tag') or '') == 'White' else 'a-colour'),
        })

    # clean up stale images
    for f in os.listdir(IMG_OUT):
        if f.endswith('.webp') and f not in used:
            os.remove(os.path.join(IMG_OUT, f))

    public = {
        'shop': S.get('Shop name') or 'Loom2hand', 'tagline': S.get('Tagline') or '',
        'wa': re.sub(r'\D', '', str(S.get('WhatsApp number') or '')),
        'phone': S.get('Phone (display)') or '', 'email': S.get('Email') or '',
        'location': S.get('Location') or '', 'instagram': S.get('Instagram URL') or '',
        'shipFee': num(S.get('Shipping fee (₹)')), 'freeAbove': num(S.get('Free shipping above (₹)')),
        'minOrder': num(S.get('Minimum order (₹)')), 'lowStock': low,
    }
    ver = time.strftime('%Y%m%d%H%M')
    data = {'v': ver, 'settings': public, 'products': products}
    js = 'window.L2H=' + json.dumps(data, ensure_ascii=False, separators=(',', ':')) + ';\n'
    with open(os.path.join(SITE, 'assets', 'products.js'), 'w', encoding='utf-8') as fh:
        fh.write(js)

    # cache-bust asset links in index.html
    idx = os.path.join(SITE, 'index.html')
    if os.path.exists(idx):
        html = open(idx, encoding='utf-8').read()
        html = re.sub(r'(assets/(?:products|app)\.js|assets/style\.css)\?v=[\w]*', lambda m: f'{m.group(1)}?v={ver}', html)
        open(idx, 'w', encoding='utf-8').write(html)

    total = sum(os.path.getsize(os.path.join(IMG_OUT, f)) for f in os.listdir(IMG_OUT))
    print(f'Built {len(products)} products, {len(used)} images ({total/1024/1024:.1f} MB), version {ver}')
    for w in warnings:
        print('  ! ' + w)
    if public['wa'] in ('', '919999999999'):
        print('  ! Set your real WhatsApp number on the Settings sheet before going live.')


if __name__ == '__main__':
    main()
